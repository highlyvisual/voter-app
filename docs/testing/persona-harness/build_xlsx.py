import json,re
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
from personas import P
d=json.load(open('merged.json'))
F='Arial'; H=Font(name=F,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='3D4A5C'); N=Font(name=F,size=10)
wrap=Alignment(wrap_text=True,vertical='top'); thin=Side(style='thin',color='D0D0D0')
def sheet(ws,headers,rows,widths):
    ws.append(headers)
    for c in ws[1]: c.font=H; c.fill=HF; c.alignment=wrap
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font=N; c.alignment=wrap; c.border=Border(bottom=thin)
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes='A2'
wb=Workbook()
# Summary
ws=wb.active; ws.title='Summary'
ws['A1']="What's It To Me? — 30 test people, run against whatsittome.org on 28 Sep 2026"; ws['A1'].font=Font(name=F,bold=True,size=13)
ws['A2']="All people are fictional. Postcodes are real. Journeys driven by an automated browser through the 'Start with you' questions."; ws['A2'].font=Font(name=F,italic=True,size=10)
labels=[("People tested","=COUNTA(Results!A2:A31)"),
 ("Reached the right ballot or the right 'no election' page","=COUNTIF(Results!I2:I31,\"Yes\")"),
 ("Hit at least one problem on the way","=COUNTIF(Results!K2:K31,\"?*\")"),
 ("Findings, all","=COUNTA(Findings!A2:A30)"),
 ("Findings rated High","=COUNTIF(Findings!C2:C30,\"High\")"),
 ("Findings rated Medium","=COUNTIF(Findings!C2:C30,\"Medium\")"),
 ("Ballots whose candidate count matched Democracy Club","=COUNTIF(Ballots!G2:G30,\"Yes\")"),
 ("Ballots checked","=COUNTA(Ballots!A2:A30)")]
for i,(l,f) in enumerate(labels,4):
    ws.cell(i,1,l).font=N; c=ws.cell(i,2,f); c.font=Font(name=F,bold=True)
ws.column_dimensions['A'].width=60; ws.column_dimensions['B'].width=12
# People
EXTRA={14:["First pass: 'Next' stayed greyed out after the postcode was entered before the page loaded"],18:["First pass: 'Next' stayed greyed out after the postcode was entered before the page loaded"],27:["First pass: 'Next' stayed greyed out after the postcode was entered before the page loaded"],
 1:["Money box shows +£4,186 Universal Credit for a single full-time student (generally not eligible)"],7:["Money box shows +£9,686 Universal Credit for a visa holder (most visas: no recourse to public funds)","'Earnings' line shows -£180 with no income"],
 5:["'Earnings' line shows £3,372 though not working"],9:["No message that 16-17s can't vote in English council elections","Household summary hidden because one question was skipped"],
 26:["Page says '0 candidates are standing' — candidates not yet confirmed"],25:["Page says '1 candidate is standing … 1 candidates' — candidates not yet confirmed"],17:["NI council row: 'ward not listed', no councillors"],
 28:["Page says '0 candidates are standing' — candidates not yet confirmed"],4:["Focus drops to the page after every answer; nothing announced to a screen reader"]}
rows=[]
for p in P:
    y=p['yn']; yn=lambda k: {True:'Yes',False:'No',None:'(skipped)'}[y[k]]
    rows.append([p['id'],p['name'],p['about'],p['postcode']+(f" (typed '{p['typo']}' first)" if p.get('typo') else ''),p['age'] or '(skipped)',p['household'] or '(skipped)',p['children'] or '(skipped)',p['tenure'] or '(skipped)',p['income'] or '(skipped)',p['work'] or '(skipped)',p['study'] or '(skipped)',yn('disability'),yn('carer'),yn('visa'),yn('benefits'),yn('drives'),yn('veteran'),p['access'].replace('_',' '),p['first_language'],p['confidence']])
sheet(wb.create_sheet('People'),['#','Name','Who they are','Postcode','Age','Adults','Children','Housing','Income','Work','Studying','Disability','Carer','Visa/asylum','Means-tested benefit','Drives','Veteran','How they use the web (test mode)','First language','Digital confidence'],rows,[4,10,48,16,11,16,14,18,16,20,14,9,7,9,10,7,8,18,12,11])
# Results
rows=[]
for p in P:
    r=d[str(p['id'])]; bt=r.get('ballot_text') or ''
    def money(label):
        m=re.search(re.escape(label)+r'\t([−+-]?\s?£[\d,]+|-£[\d,]+)',bt); return m.group(1).replace('−','-').replace(' ','') if m else ''
    hh=re.search(r'This household:(.*)',bt)
    got = r.get('ballot') or (r.get('result_url','').split('?')[0].replace('https://whatsittome.org','') if r.get('result_url') else r.get('outcome'))
    right = 'Yes' if ((r.get('ballot')==p['expect']) if p['expect'] else (r.get('ballot') is None and 'place' in (r.get('result_url') or ''))) else ('Blocked' if 'BLOCKED' in (r.get('outcome') or '') else 'No')
    cands=re.search(r'(\d+)\ncandidates? for',bt)
    pos=re.search(r'(\d+)\nsourced positions',bt)
    sf=re.search(r'Shown first for you\n\n(.*?)\n\nExplore',bt,re.S)
    issues=list(r.get('issues') or [])+EXTRA.get(p['id'],[])
    rows.append([p['id'],p['name'],p['postcode'],p['access'].replace('_',' '),p['expect'] or 'No election at this postcode',got,cands.group(1) if cands else '',pos.group(1) if pos else '',right,
        hh.group(1).strip() if hh else ('(no household summary shown)' if bt else ''),'; '.join(issues),
        (sf.group(1).replace('\n',' / ') if sf else ''),money('Universal Credit received'),money('Net income after tax and benefits'),
        len([v for k in ('axe_wizard_step9','axe_result','axe_ballot') for v in (r.get(k) or [])]),r.get('submit_s')])
sheet(wb.create_sheet('Results'),['#','Name','Postcode','Test mode','Expected','What the site showed','Candidates shown','Sourced positions on ballot','Right place?','Household summary the site showed','Problems hit on the way','Topics shown first (and why)','Universal Credit in money box','Net income in money box','Automated accessibility failures (axe, 3 pages)','Seconds from last click to results'],rows,[4,10,12,13,34,40,10,10,9,50,50,50,12,12,12,10])
# Findings
FIND=json.load(open('findings.json'))
sheet(wb.create_sheet('Findings'),['#','Area','Severity','What happened','Who it hit','How we know','Suggested fix'],[[i+1]+f for i,f in enumerate(FIND)],[4,14,9,55,30,45,45])
# Ballots
st=json.load(open('ballot_stats.json')); dc=json.load(open('dc_counts.json'))
sheet(wb.create_sheet('Ballots'),['Ballot','Candidates on site','Sourced positions','Candidates with anything published','Named sources','Democracy Club count','Counts match?','Candidates confirmed (DC)'],
      [[b,int(v['cands']),int(v['positions']),int(v['withpub']),int(v['sources']),dc[b][0],f'=IF(B{i}=F{i},"Yes","No")',dc[b][1]] for i,(b,v) in enumerate(st,2)],[62,11,11,14,10,12,10,13])
wb.save('/mnt/user-data/outputs/persona-test-results-2026-09-28.xlsx')
